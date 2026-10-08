import sys,json,re,hashlib,zipfile
from datetime import datetime,date
import openpyxl
def num(v):return v if isinstance(v,(int,float)) and not isinstance(v,bool) else None
def val(v):return v.isoformat() if isinstance(v,(date,datetime)) else v
def day(v):
    if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
    m=re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})',str(v).strip())
    if m:
        d,mo,y=map(int,m.groups());y=y-543 if y>2400 else y
        try:return date(y,mo,d).isoformat()
        except ValueError:pass
    return None
def parse(p):
    with zipfile.ZipFile(p) as z:
        if sum(i.file_size for i in z.infolist())>40000000 or len(z.infolist())>2000:raise ValueError('ไฟล์ใหญ่เกินขอบเขตที่รองรับ')
    w=openpyxl.load_workbook(p,data_only=False,keep_links=False);cw=openpyxl.load_workbook(p,data_only=True,keep_links=False)
    if 'รายการเว็บ' in w and w['รายการเว็บ']['A1'].value=='TAISIAM_LEDGER_V1':return parse_web(w,cw,p)
    sheets=[s for s in w if s['B1'].value=='บัญชีขายรายวัน' and s['C3'].value=='เปิด' and s['D3'].value=='ปิด']
    if len(sheets)!=1:raise ValueError('รองรับรูปแบบบัญชีขายรายวันที่ส่งมาเท่านั้น ต้องมีหนึ่งชีตที่หัวตารางตรงกัน')
    s=sheets[0];c=cw[s.title]
    if s.max_row>10000 or s.max_column>100:raise ValueError('ขอบเขตชีตใหญ่เกินกำหนด')
    groups=[];g=None;last=None;empty=0
    for r in range(8,s.max_row+1):
        if s.cell(r,1).value is not None:
            g={'sourceRow':r,'rawDate':str(val(s.cell(r,1).value)),'date':day(s.cell(r,1).value),'lines':[],'issues':[],'legacyDaily':num(c.cell(r,30).value)};groups.append(g);last=None
        if g is None:continue
        pump=s.cell(r,2).value
        if pump not in ('A','B','C','D',None):continue
        split=pump is None and last is not None and any(num(s.cell(r,i).value) is not None for i in (11,12,13))
        present=any(s.cell(r,i).value is not None and s.cell(r,i).data_type!='f' for i in (3,4,5,18,20,21,22,23,24))
        if not present and not split:
            if pump:empty+=1
            continue
        if pump:last=pump
        pump=pump or last
        if not pump:continue
        col={'A':11,'B':11,'C':12,'D':13}[pump];n=lambda i:num(c.cell(r,i).value)
        l={'row':r,'pump':pump,'fuel':{'A':'ดีเซล','B':'ดีเซล','C':'95','D':'91'}[pump],'split':split,'raw':{openpyxl.utils.get_column_letter(i):{'value':val(s.cell(r,i).value),'cached':val(c.cell(r,i).value)} for i in range(1,31) if s.cell(r,i).value is not None}}
        for k,i in {'opening':3,'closing':4,'liters':col,'received':5,'purchasePrice':6,'cost':14,'price':16,'legacySales':17,'legacyR':18,'legacyS':19,'discount':20,'credit':21,'legacyV':22,'legacyW':23,'debtPaid':24}.items():l[k]=n(i)
        g['lines'].append(l)
        if l['liters'] is None:g['issues'].append(f'แถว {r}: สูตรไม่มีค่าลิตรล่าสุด ให้เปิดและบันทึกใน Excel ก่อน')
        if not split and all(l[k] is not None for k in ('opening','closing','liters')) and abs(l['closing']-l['opening']-l['liters'])>.001:g['issues'].append(f'แถว {r}: ลิตรตามมิเตอร์ต่างจากลิตรลงบัญชี อาจมีการแยกต้นทุน')
        if all(l[k] is not None for k in ('liters','price','legacySales')) and abs(l['liters']*l['price']-l['legacySales'])>.01:g['issues'].append(f'Q{r}: ยอดสูตรเดิมต่างจากลิตร × ราคาขาย')
        if any(l[k] not in (None,0) for k in ('legacyR','legacyV','legacyW')):g['issues'].append(f'แถว {r}: ยอดวัดจริง/เกิน/หายเป็นลิตร ไม่หักจากเงินรับเหมือนสูตรเดิม')
    groups=[g for g in groups if g['lines']];seen=set();prev=None
    for g in groups:
        d=g['date']
        if not d:g['issues'].append('วันที่ไม่ถูกต้อง ต้องระบุวันที่จริง')
        if d in seen:g['issues'].append('วันที่ซ้ำในไฟล์ ต้องตรวจรอบการขาย')
        if d and prev and d<prev:g['issues'].append('วันที่ย้อนกลับเมื่อเทียบกับบล็อกก่อนหน้า')
        if d:seen.add(d);prev=d
        g['issues'].append('ยืนยันวันที่กับเอกสารจริง: ไฟล์มีวันที่ Excel และข้อความปะปนกัน')
        g['sales']=round(sum((l['liters'] or 0)*(l['price'] or 0) for l in g['lines']),2);g['liters']=sum(l['liters'] or 0 for l in g['lines'])
    return {'format':'legacy','sheet':s.title,'sha256':hashlib.sha256(open(p,'rb').read()).hexdigest(),'groups':groups,'emptyPumpRows':empty,'sourceSummary':{c.cell(138,i).coordinate:val(c.cell(138,i).value) for i in range(1,31) if c.cell(138,i).value is not None},'opening':[{'fuel':f,'liters':c[a].value,'cell':a} for f,a in [('ดีเซล','E4'),('95','E6'),('91','E7')]],'warnings':['ระบุสาขาก่อนยืนยัน: ผู้ใช้ยืนยันไฟล์ต้นแบบนี้เป็นสาขา TAI SIAM Tachileik','รักษาแถวแยกต้นทุนและไม่รวมยอดมิเตอร์ซ้ำ','ยอดวัดจริง/เกิน/หายเป็นลิตร ไม่ปนกับยอดเงิน; ยังไม่ทราบการผูกหัวจ่ายกับถังจริง','สูตรเดิมใช้ค่าที่ Excel บันทึกล่าสุด เก็บต้นฉบับไว้สำหรับเทียบ']}
def parse_web(w,cw,p):
    s=w['รายการเว็บ'];groups={}
    if s.max_row>10000:raise ValueError('จำนวนแถวเกินกำหนด')
    for r in range(5,s.max_row+1):
        if not s.cell(r,6).value:continue
        if any(s.cell(r,i).data_type=='f' for i in [1,2,3,4,5,6,7,8,9,10,12,13,15,16,17,18,19,20,21]):raise ValueError(f'แถว {r}: ช่องกรอกต้องเป็นค่า ไม่ใช่สูตร เพื่อป้องกันค่าคำนวณค้างจาก Excel')
        n=lambda i:num(cw[s.title].cell(r,i).value)
        rid=str(s.cell(r,3).value or '');d=day(s.cell(r,1).value);branch=str(s.cell(r,2).value or '').strip();key=(rid,branch,d)
        if key not in groups:groups[key]={'recordId':'' if rid.startswith('review:') else rid,'branch':branch,'date':d,'rawDate':str(val(s.cell(r,1).value)),'sourceRow':r,'lines':[],'issues':[]}
        g=groups[key];opening,closing=n(8),n(9);override=n(10);liters=override if override is not None else closing-opening if closing is not None and opening is not None else None
        l={'row':r,'lineId':str(s.cell(r,4).value or r),'pump':str(s.cell(r,6).value),'fuel':str(s.cell(r,7).value),'opening':opening,'closing':closing,'liters':liters,'split':s.cell(r,5).value=='แยกต้นทุน','raw':{openpyxl.utils.get_column_letter(i):{'value':val(s.cell(r,i).value),'cached':val(cw[s.title].cell(r,i).value)} for i in range(1,25) if s.cell(r,i).value is not None}}
        for k,i in {'cost':12,'price':13,'discount':15,'credit':16,'debtPaid':17,'received':18,'legacyR':19,'legacyV':20,'legacyW':21}.items():l[k]=n(i)
        g['lines'].append(l)
        if liters is None:g['issues'].append(f'แถว {r}: ไม่มีมิเตอร์หรือยอดลิตรจัดสรร')
        if override is not None:g['issues'].append(f'แถว {r}: ใช้ลิตรจัดสรร {override:g} แทนผลต่างมิเตอร์ กรุณาตรวจ')
    result=list(groups.values())
    for g in result:
        g['sales']=round(sum((l['liters'] or 0)*(l['price'] or 0) for l in g['lines']),2);g['liters']=sum(l['liters'] or 0 for l in g['lines'])
    return {'format':'web','sheet':s.title,'sha256':hashlib.sha256(open(p,'rb').read()).hexdigest(),'groups':result,'emptyPumpRows':0,'sourceSummary':{},'opening':[],'warnings':['นำกลับจาก Excel ที่เว็บส่งออก คำนวณลิตรและยอดขายจากช่องกรอกใหม่','ตรวจชื่อสาขา วันที่ รหัสรายการและการเปลี่ยนแปลงก่อนยืนยัน','ไม่มีการลบข้อมูลบนเว็บจากการลบแถวใน Excel']}
if __name__=='__main__':
    try:print(json.dumps(parse(sys.argv[1]),ensure_ascii=False,allow_nan=False))
    except Exception as e:print(json.dumps({'error':str(e)},ensure_ascii=False));sys.exit(1)
